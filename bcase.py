import streamlit as st
import pandas as pd
import numpy as np
import io
import matplotlib.pyplot as plt
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as OpenpyxlImage

# --- 1. Advanced Financial Formatting & Red/Green Engine ---
def format_fin(val, row_name=""):
    if pd.isna(val) or (isinstance(val, str) and val.strip() in ["-", ""]): return "-"
    try:
        if isinstance(val, str): val = val.replace(',', '').replace(' ', '').replace('%', '')
        num = float(val)
        if any(x in str(row_name).upper() for x in ["%", "RATE", "OCC"]):
            return f"{num:.1%}"
        return f"{num:,.1f}"
    except: return str(val)

def render_table(df_raw, item_col="P&L Line Item", variance_cols=["Variance"]):
    df_disp = df_raw.copy()
    for col in df_disp.columns:
        if col != item_col:
            df_disp[col] = df_raw.apply(lambda row: format_fin(row[col], row[item_col]), axis=1)
            
    def style_cells(disp_row):
        styles = []
        idx = disp_row.name
        row_name_raw = str(disp_row[item_col])
        row_name = row_name_raw.strip()
        
        is_bold_row = row_name in ["Business volume HT", "Total revenue", "Variable margin", "GOP Total"]
        is_sub_row = row_name_raw.startswith("  -")
        
        for col in df_disp.columns:
            # 🌟 Updated: Base styling with center alignment
            cell_style = 'background-color: #F8F9FA !important; color: #00204A !important; text-align: center !important;'
            
            if col == item_col:
                cell_style += ' font-weight: bold; text-align: left !important;' # Keep items left-aligned for readability
                if is_bold_row: 
                    cell_style = 'background-color: #E6F2FF !important; color: #00204A !important; font-weight: 900; border-top: 2px solid #00204A; text-align: left !important;'
                elif is_sub_row:
                    cell_style = 'background-color: #F8F9FA !important; color: #6C757D !important; font-style: italic; font-weight: normal; padding-left: 20px; text-align: left !important;'
            else:
                if is_bold_row: 
                    cell_style += ' font-weight: 900; border-top: 2px solid #00204A; background-color: #E6F2FF !important;'
                elif is_sub_row:
                    cell_style += ' color: #6C757D !important; font-style: italic;'
                
                # Variance Heat-map
                if any(v_col in col for v_col in variance_cols):
                    val = df_raw.iloc[idx][col]
                    try:
                        v = float(str(val).replace(',', '').replace('%', '')) if not pd.isna(val) else 0
                        if abs(v) > 0.001:
                            if v > 0: cell_style += ' color: #008000 !important; font-weight: bold;'
                            else: cell_style += ' color: #FF0000 !important; font-weight: bold;'
                    except: pass
            styles.append(cell_style)
        return styles

    styled_df = df_disp.style.apply(style_cells, axis=1)
    
    # 🌟 Updated: Center-aligned headers and increased font size (18px)
    header_style = [
        {'selector': 'th', 'props': [
            ('background-color', '#00204A !important'), 
            ('color', '#FFFFFF !important'), 
            ('font-weight', 'bold !important'), 
            ('text-align', 'center !important'), 
            ('font-size', '18px !important'), 
            ('height', '45px !important'), 
            ('vertical-align', 'middle !important')
        ]},
        {'selector': 'td', 'props': [
            ('border-bottom', '1px solid #E0E0E0'), 
            ('height', '38px !important'), 
            ('vertical-align', 'middle !important')
        ]}
    ]
    styled_df.set_table_styles(header_style)
    st.table(styled_df)

# --- Season Dictionary ---
MONTH_MAP = {
    "Winter": ["November", "December", "January", "February", "March", "April"],
    "Summer": ["May", "June", "July", "August", "September", "October"],
    "S1": ["January", "February", "March", "April", "May", "June"],
    "S2": ["July", "August", "September", "October", "November", "December"],
    "Full Year": ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]
}

# --- 2. Parsing Engine ---
def get_bench_data(df, resort, year, months_list):
    sub = df[(df['Resort'] == resort) & (df['Year'] == str(year)) & (df['Month'].isin(months_list))]
    
    def fetch_val(line_item_name):
        match = sub[sub['Line_Item'] == line_item_name]
        if match.empty: return 0.0
        s = match['Amount_kRMB'].astype(str).str.replace(',', '').str.replace(' ', '').replace('-', '0')
        return pd.to_numeric(s, errors='coerce').sum()

    capa = fetch_val("Capacity")
    hn = fetch_val("Nbr of hotel days on site") + fetch_val("HN Operation area country") + fetch_val("HN Operation area zone")
    bv_ttc = fetch_val("Business Volume TTC")
    
    li_match = sub[sub['Line_Item'].str.contains('Local income|VSL/VRL|Operating Revenues|Local Operating', case=False, na=False)]
    li = pd.to_numeric(li_match['Amount_kRMB'].astype(str).str.replace(',', '').replace(' ', '').replace('-', '0'), errors='coerce').sum()
    
    vc_fb = fetch_val("Purchase Food & Beverage GM")
    vc_ski = fetch_val("Purchase of ski lift & ski pass")
    vc_total = fetch_val("Total Variable Charges")
    fc_total = fetch_val("Total Fixed Charges")
    
    etp_gof = fetch_val("Nbr of ETP seasonal GO  - Contrat SEC")
    etp_gol = fetch_val("Nbr of ETP seasonal GO  - Contrat LOC")
    etp_ge = fetch_val("Nbr of ETP seasonal GE")
    
    sal_gof = fetch_val("Salaries & charges - Seasonal GO - Contract sector")
    sal_gol = fetch_val("Salaries & charges - Seasonal GO - Contract local") + fetch_val("Salaries & charges - Permanent GO")
    sal_ge = fetch_val("Salaries & charges - Seasonal GE") + fetch_val("Salaries & charges - GE Intern")
    
    go_fee = abs(fetch_val("GO fee"))
    m_fee_raw = abs(fetch_val("Management fees")) + abs(fetch_val("Management fees to Club Med"))
    if m_fee_raw == 0: m_fee_raw = ((bv_ttc * 0.7)/1.06 + li) * 0.02

    vc_other = vc_total - vc_fb - vc_ski + m_fee_raw
    fc_other = fc_total - sal_gof - sal_gol - sal_ge + go_fee

    return {
        "Capa": capa, "HN": hn, "BV": bv_ttc, "LI": li, 
        "VC_FB": vc_fb, "VC_Ski": vc_ski, "VC_Other": vc_other,
        "ETP_GOF": etp_gof, "ETP_GOL": etp_gol, "ETP_GE": etp_ge,
        "Sal_GOF": sal_gof, "Sal_GOL": sal_gol, "Sal_GE": sal_ge, 
        "FC_Other": fc_other,
        "GO_Fee_Raw": go_fee,
        "N_months": len(months_list)
    }

def merge_periods(d1, d2):
    res = {}
    for k in ['Capa', 'HN', 'BV', 'LI', 'VC_FB', 'VC_Ski', 'VC_Other', 'FC_Other', 'Sal_GOF', 'Sal_GOL', 'Sal_GE', 'GO_Fee_Raw', 'N_months']:
        res[k] = d1.get(k,0) + d2.get(k,0)
    res['ETP_GOF'] = (d1['ETP_GOF']*d1['N_months'] + d2['ETP_GOF']*d2['N_months']) / res['N_months'] if res['N_months'] else 0
    res['ETP_GOL'] = (d1['ETP_GOL']*d1['N_months'] + d2['ETP_GOL']*d2['N_months']) / res['N_months'] if res['N_months'] else 0
    res['ETP_GE'] = (d1['ETP_GE']*d1['N_months'] + d2['ETP_GE']*d2['N_months']) / res['N_months'] if res['N_months'] else 0
    return res

# --- 3. P&L Flow Logic ---
def get_pl_flow(d, detailed=False, is_benchmark=False):
    hn = d['HN']
    capa = d.get('Capa', 0)
    bv_ttc = d['BV']
    adr = (bv_ttc * 1000 / hn) if hn else 0
    occ = (hn / capa) if capa != 0 else 0
    
    tot_bv_ttc = bv_ttc
    cm_bv_ttc = bv_ttc * 0.3
    own_bv_ttc = bv_ttc * 0.7
    tot_bv_ht = tot_bv_ttc / 1.06
    own_bv_ht = own_bv_ttc / 1.06
    cm_bv_ht = cm_bv_ttc
    
    tot_vat = -(own_bv_ht * 0.06)
    own_vat = tot_vat
    cm_vat = 0
    
    li = d['LI']
    tot_rev = tot_bv_ht + li
    own_rev = own_bv_ht + li
    cm_rev = cm_bv_ttc
    
    m_fee = own_rev * 0.02
    go_fee = d['GO_Fee_Raw'] if is_benchmark else abs(d.get('Sal_GOF', 0) + d.get('Sal_GOL', 0)) * 0.20
    
    vc_base = d.get('VC_FB', 0) + d.get('VC_Ski', 0) + d.get('VC_Other', 0)
    fc_base = d.get('Sal_GOF', 0) + d.get('Sal_GOL', 0) + d.get('Sal_GE', 0) + d.get('FC_Other', 0)
    
    tot_vc = vc_base
    own_vc = vc_base - m_fee
    cm_vc = m_fee
    tot_fc = fc_base
    own_fc = fc_base - go_fee
    cm_fc = go_fee
    
    own_var_margin = own_rev + own_vc
    cm_var_margin = cm_rev + cm_vc
    tot_var_margin = tot_rev + tot_vc
    
    own_gop = own_var_margin + own_fc
    cm_gop = cm_var_margin + cm_fc
    tot_gop = tot_var_margin + tot_fc

    rows = [
        ["ADR (RMB)", adr, np.nan, adr],
        ["HN sold", hn, np.nan, hn],
        ["OCC %", occ, np.nan, occ],
        ["Business Volume TTC", tot_bv_ttc, cm_bv_ttc, own_bv_ttc],
        ["VAT on outside turnover", tot_vat, cm_vat, own_vat],
        ["Business volume HT", tot_bv_ht, cm_bv_ht, own_bv_ht],
        ["Local income HT", li, 0, li],
        ["Total revenue", tot_rev, cm_rev, own_rev],
        ["Total variable cost", tot_vc, cm_vc, own_vc],
    ]
    if detailed:
        rows += [
            ["  - Purchase Food & Beverage GM", d.get('VC_FB',0), 0, d.get('VC_FB',0)],
            ["  - Purchase of ski lift & ski pass", d.get('VC_Ski',0), 0, d.get('VC_Ski',0)],
        ]
    rows += [
        ["  - Management fees to Club Med", 0, m_fee, -m_fee], 
        ["Variable margin", tot_var_margin, cm_var_margin, own_var_margin],
        ["Total fixed cost", tot_fc, cm_fc, own_fc]
    ]
    if detailed:
        rows += [
            ["  - Staff expense Foreign GO", d.get('Sal_GOF',0), 0, d.get('Sal_GOF',0)],
            ["  - Staff expense Local GO", d.get('Sal_GOL',0), 0, d.get('Sal_GOL',0)],
            ["  - Staff expense GE", d.get('Sal_GE',0), 0, d.get('Sal_GE',0)]
        ]
    rows += [
        ["  - GO fees", 0, go_fee, -go_fee], 
        ["GOP Total", tot_gop, cm_gop, own_gop],
        ["GOP Margin %", tot_gop/tot_rev if tot_rev!=0 else 0, cm_gop/cm_rev if cm_rev!=0 else 0, own_gop/own_rev if own_rev!=0 else 0]
    ]
    return pd.DataFrame(rows, columns=["P&L Line Item", "Total", "CM Portion", "Owner Portion"])

def filter_pl_view(df_flow, view_mode):
    col_name = "Total"
    if view_mode == "CM View": col_name = "CM Portion"
    elif view_mode == "Owner View": col_name = "Owner Portion"
    return df_flow[["P&L Line Item", col_name]]

# --- 4. Excel Styling ---
def style_excel_sheet(ws, df):
    header_fill = PatternFill(start_color="00204A", end_color="00204A", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    bold_rows = ["Business volume HT", "Total revenue", "Variable margin", "GOP Total"]
    bold_fill = PatternFill(start_color="E6F2FF", end_color="E6F2FF", fill_type="solid")
    
    for r_idx, row in enumerate(df.values, start=2):
        row_name_raw = str(row[0])
        row_name = row_name_raw.strip()
        is_bold = row_name in bold_rows
        is_sub = row_name_raw.startswith("  -")
        
        for c_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx)
            if isinstance(val, (int, float)):
                if any(x in str(row_name).upper() for x in ["%", "RATE", "OCC"]): cell.number_format = '0.0%'
                else: cell.number_format = '#,##0.1'
            
            if c_idx == 1:
                if is_bold:
                    cell.font = Font(bold=True, color="00204A")
                    cell.fill = bold_fill
                elif is_sub:
                    cell.font = Font(color="6C757D", italic=True)
                    cell.alignment = Alignment(indent=2)
                else:
                    cell.font = Font(bold=True, color="00204A")
            else:
                if is_bold:
                    cell.font = Font(bold=True, color="00204A")
                    cell.fill = bold_fill

    ws.column_dimensions['A'].width = 38
    for col in range(2, len(df.columns)+1): ws.column_dimensions[get_column_letter(col)].width = 15

# --- Header Renderer ---
def render_module_header(title, icon=""):
    st.markdown(f"""
    <div style="background: linear-gradient(135deg, #1D263B 0%, #2A3650 100%); padding: 18px 20px; border-radius: 10px; text-align: center; margin-bottom: 25px; box-shadow: 0 4px 12px rgba(0,0,0,0.1);">
        <h2 style="margin: 0; color: #FFFFFF; font-family: 'Playfair Display', serif; font-size: 1.8rem; letter-spacing: 1px;">
            <span style="color: #A64B35; margin-right: 10px;">{icon}</span>{title}
        </h2>
    </div>
    """, unsafe_allow_html=True)

# --- Streamlit UI Config ---
st.set_page_config(layout="wide", page_title="B-Case Decision Engine Pro V40")

# --- CSS Global Styles ---
st.markdown("""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Playfair+Display:wght@600;700&family=Inter:wght@300;400;500;600&display=swap');
    
    h1, h2, h3 { font-family: 'Playfair Display', serif !important; color: #1D263B; }
    .stDataFrame { border-radius: 8px; overflow: hidden; box-shadow: 0 4px 15px rgba(0,0,0,0.03); }
    .stSidebar { background-color: #F8F9FA !important; border-right: 1px solid #EAECEF; }
    
    div[data-testid="stTabs"] button {
        flex: 1; font-family: 'Inter', sans-serif; font-weight: 600; color: #6C757D;
        background-color: #F8F9FA; border: 1px solid #EAECEF; border-bottom: 2px solid transparent;
        border-radius: 8px 8px 0 0; padding: 14px 10px; margin: 0 4px;
    }
    div[data-testid="stTabs"] button[aria-selected="true"] {
        background-color: #FFFFFF; color: #1D263B; border-bottom: 4px solid #A64B35;
    }
</style>
""", unsafe_allow_html=True)

# --- Authentication ---
if "authenticated" not in st.session_state:
    st.session_state["authenticated"] = False

if not st.session_state["authenticated"]:
    st.title("🔐 Club Med B-Case Tool")
    password = st.text_input("Access Password", type="password")
    if st.button("Login"):
        if password == "CM2026Daniel": 
            st.session_state["authenticated"] = True
            st.rerun()
        else:
            st.error("Incorrect password. Please contact Daniel.")
    st.stop() 

st.sidebar.header("📂 Digital Ledger")
uploaded_file = st.sidebar.file_uploader("Upload RAW DATA.csv", type="csv")

if uploaded_file:
    df = pd.read_csv(uploaded_file)
    df.columns = [c.strip() for c in df.columns]
    
    for col in ['Resort', 'Year', 'Month', 'Line_Item']:
        if col in df.columns:
            df[col] = df[col].astype(str).str.strip()
            if col == 'Year':
                df[col] = df[col].apply(lambda x: x[:-2] if x.endswith('.0') else x)
    
    resorts = sorted(df['Resort'].unique())
    sel_resort = st.sidebar.selectbox("Select Resort", resorts)
    
    available_years = df[df['Resort'] == sel_resort]['Year'].astype(str).unique()
    available_years = sorted(available_years, reverse=True)
    sel_year = st.sidebar.selectbox("Select Year", available_years)

    # --- Step Navigation ---
    tab1, tab2, tab3, tab4, tab5, tab6 = st.tabs([
        "STEP 1: Benchmark Dashboard", 
        "STEP 2: Business Forecast", 
        "STEP 3: Variance Analysis", 
        "STEP 4: Cost Fine-Tuning", 
        "STEP 5: Stress Test", 
        "STEP 6: 10-Year Simulation"
    ])

    m1_mode = st.sidebar.radio("Global Dimension View:", ["By Season (Winter/Summer)", "By Semester (S1/S2)"])
    m1_cols = ["Winter", "Summer", "Full Year"] if "Season" in m1_mode else ["S1", "S2", "Full Year"]
    p1_name, p2_name = m1_cols[0], m1_cols[1]

    # --- STEP 1: Benchmark ---
    with tab1:
        render_module_header(f"Module 1: {sel_resort} {sel_year} Benchmark Board", "📊")
        st.subheader("1.1 Core Operational KPIs")
        ops_res = []
        for c in m1_cols:
            d = get_bench_data(df, sel_resort, sel_year, MONTH_MAP[c])
            ops_res.append([d['Capa'], d['HN'], d['BV'], (d['BV']*1000)/d['HN'] if d['HN']!=0 else 0, d['HN']/d['Capa'] if d['Capa']!=0 else 0])
        ops_df = pd.DataFrame(ops_res, index=m1_cols, columns=["Capacity", "HN sold", "BV total (kRMB)", "ADR (RMB)", "OCC %"]).T
        render_table(ops_df.reset_index(), item_col="index")

        st.subheader("1.2 P&L Flow & Distribution")
        m1_sel = st.selectbox("Select P&L Interval", m1_cols)
        render_table(get_pl_flow(get_bench_data(df, sel_resort, sel_year, MONTH_MAP[m1_sel]), detailed=False, is_benchmark=True))
        st.markdown("<div style='text-align: right; color: #A64B35; font-weight: bold;'>Click Step 2 Above ➡️</div>", unsafe_allow_html=True)

    # --- STEP 2: Forecast ---
    with tab2:
        render_module_header("Module 2: Operational Entry (V1 Forecast)", "📥")
        x_ratio = st.number_input("Parameter X (Local Income as % of BV TTC)", value=10.0) / 100
        d_ref_full = get_bench_data(df, sel_resort, sel_year, MONTH_MAP["Full Year"])
        capa_default = d_ref_full['Capa'] / 12 if d_ref_full['Capa'] else 23000.0
        input_df = pd.DataFrame({"Month": MONTH_MAP["Full Year"], "Capacity": [capa_default]*12, "HN Sold": [8000.0]*12, "ADR (RMB)": [1200.0]*12})
        e_df = st.data_editor(input_df, hide_index=True)
        e_df['BV'] = (e_df['HN Sold'] * e_df['ADR (RMB)']) / 1000

        def get_v1_stats(m_list):
            ref = get_bench_data(df, sel_resort, sel_year, m_list)
            sub = e_df[e_df['Month'].isin(m_list)]
            hn, bv, capa = sub['HN Sold'].sum(), sub['BV'].sum(), sub['Capacity'].sum()
            v1 = ref.copy()
            v1['HN'], v1['BV'], v1['Capa'], v1['LI'] = hn, bv, capa, bv * x_ratio
            v1['VC_FB'] = (ref['VC_FB']/ref['HN'] * hn) if ref['HN']!=0 else 0
            v1['VC_Ski'] = (ref['VC_Ski']/ref['HN'] * hn) if ref['HN']!=0 else 0
            v1['VC_Other'] = (ref['VC_Other']/ref['HN'] * hn) if ref['HN']!=0 else 0
            return v1

        v1_p1, v1_p2 = get_v1_stats(MONTH_MAP[p1_name]), get_v1_stats(MONTH_MAP[p2_name])
        v1_full = merge_periods(v1_p1, v1_p2)
        st.markdown("<div style='text-align: right; color: #A64B35; font-weight: bold;'>Click Step 3 Above ➡️</div>", unsafe_allow_html=True)

    # --- STEP 3: Variance ---
    with tab3:
        render_module_header("Module 3: V1 vs Benchmark Variance Analysis", "📉")
        c3_1, c3_2 = st.columns(2)
        m3_sel = c3_1.selectbox("Comparison Interval", m1_cols, key="m3")
        m3_view = c3_2.radio("Perspective:", ["Total View", "CM View", "Owner View"], horizontal=True, key="m3_v")
        d_m3_bench = get_bench_data(df, sel_resort, sel_year, MONTH_MAP[m3_sel])
        d_m3_v1 = v1_full if m3_sel == "Full Year" else (v1_p1 if m3_sel == p1_name else v1_p2)
        f3_b, f3_v1 = get_pl_flow(d_m3_bench, is_benchmark=True), get_pl_flow(d_m3_v1)
        df3_b_f, df3_v1_f = filter_pl_view(f3_b, m3_view), filter_pl_view(f3_v1, m3_view)
        df_m3 = pd.DataFrame({"P&L Line Item": f3_b["P&L Line Item"], "Bench": df3_b_f.iloc[:, 1], "V1 Forecast": df3_v1_f.iloc[:, 1]})
        df_m3["Variance"] = df_m3["V1 Forecast"] - df_m3["Bench"]
        render_table(df_m3, variance_cols=["Variance"])
        st.markdown("<div style='text-align: right; color: #A64B35; font-weight: bold;'>Click Step 4 Above ➡️</div>", unsafe_allow_html=True)

    # --- STEP 4: Fine-Tuning ---
    with tab4:
        render_module_header(f"Module 4: Dual-Channel Cost Optimization (V2 Budget)", "⚙️")
        def render_m4_adj(p_name, v1_data):
            st.subheader(f"🛠️ {p_name} Fine-Tuning Panel")
            hn_v1 = v1_data['HN']
            fb_uc = (v1_data['VC_FB']*1000)/hn_v1 if hn_v1 else 0
            ski_uc = (v1_data['VC_Ski']*1000)/hn_v1 if hn_v1 else 0
            s_f = (v1_data['Sal_GOF']*1000)/v1_data['ETP_GOF'] if v1_data['ETP_GOF'] else 0
            s_l = (v1_data['Sal_GOL']*1000)/v1_data['ETP_GOL'] if v1_data['ETP_GOL'] else 0
            s_ge = (v1_data['Sal_GE']*1000)/v1_data['ETP_GE'] if v1_data['ETP_GE'] else 0
            
            items = ["F&B Unit (RMB/HN)", "Ski Unit (RMB/HN)", "Sal F-GO per ETP", "Sal L-GO per ETP", "Sal GE per ETP"]
            etp_bases = ["-", "-", v1_data['ETP_GOF'], v1_data['ETP_GOL'], v1_data['ETP_GE']]
            befores = [fb_uc, ski_uc, s_f, s_l, s_ge]
            rc = st.columns([3, 2, 2, 2, 2])
            for col, title in zip(rc, ["Cost Items", "Base(HN/ETP)", "Before(V1)", "Adj %", "After(V2)"]): col.markdown(f"**{title}**")
            
            afters = []
            for i in range(5):
                rc = st.columns([3, 2, 2, 2, 2])
                rc[0].write(items[i])
                rc[1].write(f"{etp_bases[i]:,.1f}" if isinstance(etp_bases[i], (int, float)) else etp_bases[i])
                rc[2].write(f"{befores[i]:,.1f}")
                adj = rc[3].number_input(" ", value=0.0, key=f"adj_{p_name}_{i}", label_visibility="collapsed")
                after = befores[i] * (1 + adj/100)
                rc[4].write(f"`{after:,.1f}`")
                afters.append(after)
            v2 = v1_data.copy()
            v2['VC_FB'] = (afters[0] * v2['HN']) / 1000
            v2['VC_Ski'] = (afters[1] * v2['HN']) / 1000
            v2['Sal_GOF'] = (afters[2] * v2['ETP_GOF']) / 1000
            v2['Sal_GOL'] = (afters[3] * v2['ETP_GOL']) / 1000
            v2['Sal_GE'] = (afters[4] * v2['ETP_GE']) / 1000
            return v2
            
        c4_1, c4_2 = st.columns(2)
        with c4_1: v2_p1 = render_m4_adj(p1_name, v1_p1)
        with c4_2: v2_p2 = render_m4_adj(p2_name, v1_p2)
        v2_full = merge_periods(v2_p1, v2_p2)
        m4_view = st.radio("Perspective:", ["Total View", "CM View", "Owner View"], horizontal=True, key="m4_v")
        f4_v1, f4_v2 = get_pl_flow(v1_full, detailed=True), get_pl_flow(v2_full, detailed=True)
        df4_v1_f, df4_v2_f = filter_pl_view(f4_v1, m4_view), filter_pl_view(f4_v2, m4_view)
        df_m4 = pd.DataFrame({"P&L Line Item": f4_v1["P&L Line Item"], "V1 Base": df4_v1_f.iloc[:, 1], "V2 Adjusted": df4_v2_f.iloc[:, 1]})
        df_m4["Variance"] = df_m4["V2 Adjusted"] - df_m4["V1 Base"]
        render_table(df_m4, variance_cols=["Variance"])

    # --- STEP 5: Stress Test ---
    with tab5:
        render_module_header(f"Module 5: Dual-Channel Stress Test (V3 Stabilized Year)", "🧪")
        def render_m5_adj(p_name, v2_data):
            st.subheader(f"{p_name} Stress Test")
            adr_b = v2_data['BV']*1000/v2_data['HN'] if v2_data['HN']!=0 else 0
            st.write(f"Before (V2) -> ADR: **{adr_b:,.1f}** | HN: **{v2_data['HN']:,.0f}**")
            a_adr = st.slider(f"{p_name} ADR Fluctuation (%)", -30, 30, 0, key=f"m5a_{p_name}") / 100
            a_hn = st.slider(f"{p_name} HN Fluctuation (%)", -30, 30, 0, key=f"m5h_{p_name}") / 100
            v3 = v2_data.copy()
            v3['HN'] *= (1+a_hn)
            v3['BV'] *= (1+a_adr)*(1+a_hn)
            v3['LI'] = v3['BV'] * x_ratio
            v3['VC_FB'] *= (1+a_hn); v3['VC_Ski'] *= (1+a_hn); v3['VC_Other'] *= (1+a_hn)
            st.write(f"After (V3) -> ADR: **{adr_b*(1+a_adr):,.1f}** | HN: **{v3['HN']:,.0f}** | OCC: **{v3['HN']/v3['Capa'] if v3['Capa']!=0 else 0:.1%}**")
            return v3
        c5_1, c5_2 = st.columns(2)
        with c5_1: v3_p1 = render_m5_adj(p1_name, v2_p1)
        with c5_2: v3_p2 = render_m5_adj(p2_name, v2_p2)
        v3_full = merge_periods(v3_p1, v3_p2)
        m5_view = st.radio("Perspective:", ["Total View", "CM View", "Owner View"], horizontal=True, key="m5_v")
        f5_v2, f5_v3 = get_pl_flow(v2_full, detailed=True), get_pl_flow(v3_full, detailed=True)
        df5_v2_f, df5_v3_f = filter_pl_view(f5_v2, m5_view), filter_pl_view(f5_v3, m5_view)
        df_m5 = pd.DataFrame({"P&L Line Item": f5_v2["P&L Line Item"], "V2 Base": df5_v2_f.iloc[:, 1], "V3 Adjusted": df5_v3_f.iloc[:, 1]})
        df_m5["Variance"] = df_m5["V3 Adjusted"] - df_m5["V2 Base"]
        render_table(df_m5, variance_cols=["Variance"])

    # --- STEP 6: 10-Year Sandbox ---
    with tab6:
        render_module_header("Module 6: 10-Year P&L Dynamic Simulation", "📈")
        st.write("Set YoY Growth Rates based on V3基准.")
        yoy_df = pd.DataFrame({"Year": [f"Y{i:02}" for i in range(1, 11)], "Capa Growth %": [0.0]*10, "ADR Growth %": [-10.0, -10.0, 0.0] + [2.0]*7, "HN Growth %": [-20.0, -20.0, 0.0] + [3.0]*7, "Inflation %": [0.0]*10})
        e_yoy = st.data_editor(yoy_df, hide_index=True)
        if st.button("🚀 Generate 10-Year Full Projection & Charts"):
            m_hn, m_adr, m_inf, m_capa = [1.0]*11, [1.0]*11, [1.0]*11, [1.0]*11
            for i in range(1, 11):
                m_capa[i] = m_capa[i-1] * (1 + e_yoy.loc[i-1, 'Capa Growth %']/100) if i>1 else 1.0
                m_hn[i] = m_hn[i-1] * (1 + e_yoy.loc[i-1, 'HN Growth %']/100) if i>1 else 1.0
                m_adr[i] = m_adr[i-1] * (1 + e_yoy.loc[i-1, 'ADR Growth %']/100) if i>1 else 1.0
                m_inf[i] = m_inf[i-1] * (1 + e_yoy.loc[i-1, 'Inflation %']/100) if i>1 else 1.0
            pnl_total, pnl_cm, pnl_owner = {}, {}, {}
            chart_gop, chart_margin = [], []
            for y in range(1, 11):
                yr_str = f"Y{y:02}"
                d_y = v3_full.copy()
                d_y['Capa'] *= m_capa[y]; d_y['HN'] *= m_hn[y]
                d_y['BV'] = d_y['HN'] * (v3_full['BV']/v3_full['HN'] * m_adr[y]) if v3_full['HN']!=0 else 0
                d_y['LI'] = d_y['BV'] * x_ratio
                d_y['VC_FB'] *= m_hn[y] * m_inf[y]; d_y['VC_Ski'] *= m_hn[y] * m_inf[y]; d_y['VC_Other'] *= m_hn[y] * m_inf[y]
                d_y['FC_Other'] *= m_inf[y]; d_y['Sal_GOF'] *= m_inf[y]; d_y['Sal_GOL'] *= m_inf[y]; d_y['Sal_GE'] *= m_inf[y]
                f_y = get_pl_flow(d_y, detailed=False)
                pnl_total[yr_str], pnl_cm[yr_str], pnl_owner[yr_str] = f_y['Total'], f_y['CM Portion'], f_y['Owner Portion']
                t_gop = f_y.loc[f_y["P&L Line Item"] == "GOP Total", "Total"].values[0]
                o_margin = f_y.loc[f_y["P&L Line Item"] == "GOP Margin %", "Owner Portion"].values[0]
                chart_gop.append({"Year": yr_str, "Total GOP": t_gop})
                chart_margin.append({"Year": yr_str, "Owner Margin %": o_margin * 100})
            
            st.subheader("10-Year Profitability Trends")
            col_c1, col_c2 = st.columns(2)
            with col_c1: 
                fig1, ax1 = plt.subplots(figsize=(10, 5))
                pd.DataFrame(chart_gop).set_index("Year").plot(ax=ax1, marker='o', color='#00204A')
                ax1.set_title("10-Year GOP Trend (kRMB)"); st.pyplot(fig1)
            with col_c2:
                fig2, ax2 = plt.subplots(figsize=(10, 5))
                pd.DataFrame(chart_margin).set_index("Year").plot(ax=ax2, marker='s', color='#A64B35')
                ax2.set_title("10-Year Owner Margin % Trend"); st.pyplot(fig2)

            df_tot = pd.DataFrame(pnl_total); df_tot.insert(0, "P&L Line Item", f_y['P&L Line Item'])
            st.tabs(["Full 10-Year P&L"])[0].table(df_tot)
            
else:
    # 🌟 Welcome Hero Banner
    welcome_html = """
    <div style="padding: 5rem 2rem; text-align: center; background: linear-gradient(135deg, #1D263B 0%, #2A3650 100%); border-radius: 16px; margin-top: 1rem; box-shadow: 0 20px 40px rgba(0,0,0,0.15);">
        <div style="font-size: 4.5rem; margin-bottom: 0.5rem; color: #A64B35; font-family: serif;">Ψ</div>
        <h1 style="font-family: 'Playfair Display', serif; font-size: 3.5rem; color: #FFFFFF; margin-bottom: 1rem; letter-spacing: 1px;">Financial Decision Engine</h1>
        <p style="font-family: 'Inter', sans-serif; font-size: 1.15rem; color: #A4B6B0; max-width: 650px; margin: 0 auto; line-height: 1.6; font-weight: 300;">
            Elevate your financial modeling. Please upload your RAW DATA via the sidebar to unlock 10-year P&L simulations, real-time cost adjustments, and executive-level margin analysis.
        </p>
    </div>
    """
    st.markdown(welcome_html, unsafe_allow_html=True)

    st.markdown("<br><br>", unsafe_allow_html=True)

    # 🌟 Feature Cards (高级功能展示卡片 - 全英文版)
    c1, c2, c3 = st.columns(3)
    
    card_style = "padding: 2rem 1.5rem; background-color: #FFFFFF; border-radius: 12px; box-shadow: 0 4px 20px rgba(0,0,0,0.04); border-top: 4px solid #A64B35; height: 100%; text-align: center;"
    
    with c1:
        st.markdown(f'''
        <div style="{card_style}">
            <div style="font-size: 2.5rem; margin-bottom: 1rem;">📊</div>
            <h3 style="font-family: 'Playfair Display', serif; color: #1D263B; font-size: 1.4rem; margin-bottom: 0.5rem;">P&L Benchmarking</h3>
            <p style="color: #6c757d; font-size: 0.95rem; line-height: 1.5;">Instantly compare operational KPIs and P&L flows across seasons and semesters with boardroom-ready precision.</p>
        </div>
        ''', unsafe_allow_html=True)
        
    with c2:
        st.markdown(f'''
        <div style="{card_style}">
            <div style="font-size: 2.5rem; margin-bottom: 1rem;">⚙️</div>
            <h3 style="font-family: 'Playfair Display', serif; color: #1D263B; font-size: 1.4rem; margin-bottom: 0.5rem;">Dynamic Scenarios</h3>
            <p style="color: #6c757d; font-size: 0.95rem; line-height: 1.5;">Simulate V1 Forecasts, fine-tune V2 Budgets, and execute V3 Stress Tests in real-time.</p>
        </div>
        ''', unsafe_allow_html=True)
        
    with c3:
        st.markdown(f'''
        <div style="{card_style}">
            <div style="font-size: 2.5rem; margin-bottom: 1rem;">📈</div>
            <h3 style="font-family: 'Playfair Display', serif; color: #1D263B; font-size: 1.4rem; margin-bottom: 0.5rem;">10-Year Projection</h3>
            <p style="color: #6c757d; font-size: 0.95rem; line-height: 1.5;">Extrapolate long-term GOP trends and margins with customizable YoY growth rates and inflation modeling.</p>
        </div>
        ''', unsafe_allow_html=True)
